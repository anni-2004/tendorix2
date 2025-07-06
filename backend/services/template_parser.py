import os
import re
import json
from docx import Document
import google.generativeai as genai  # Only used to access the model


# Direct API key setup (for testing only)
genai.configure(api_key="AIzaSyChb5KD4nkdMJMAOXtG1yzU-7awgBFFx5Y")
MODEL = genai.GenerativeModel("gemini-2.0-flash-lite")

def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)

def build_prompt(template_text):
    return f"""
You are a document template parser.

Your task is to convert any given pharmaceutical or tender-based document template into a JSON schema with:
- Field metadata (ID, label, type, generative)
- A templateString using {{placeholders}} where data goes

👀 Detect dynamic fields such as:
- Underlines (_________), placeholders ([Date], [Deviation 1])
- Table headers with repeating data (e.g., Brand Name, Generic Name, Pack Size)

If you see a table-like structure, define it as:
"type": "array of objects" and include its "itemSchema"

🎯 OUTPUT FORMAT:
{{
  "name": "<Template Name>",
  "fields": [...],
  "templateString": "..."
}}

🎓 Field Types:
- string
- date
- array of objects

📌 Only return JSON. Do NOT wrap in ```json.

---

Document Template:

{template_text}

---
"""

def clean_and_parse_gemini_json(text):
    try:
        cleaned = re.sub(r"^```json\s*|\s*```$", "", text.strip(), flags=re.IGNORECASE)
        return json.loads(cleaned)
    except Exception as e:
        raise ValueError(f"❌ Failed to parse Gemini output: {e}")

def extract_schema_from_docx(docx_path):
    template_text = extract_text_from_docx(docx_path)
    prompt = build_prompt(template_text)

    response = MODEL.generate_content(prompt)

    try:
        parsed = clean_and_parse_gemini_json(response.text)
        print("✅ Extracted schema:")
        print(json.dumps(parsed, indent=2))
        return parsed
    except Exception as e:
        print("❌ Error:", e)
        print("📄 Raw Gemini output:\n", response.text)
        return None
